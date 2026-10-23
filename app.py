import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 1. Criar a pasta de trabalho e as abas
wb = openpyxl.Workbook()

# Aba de Resumo/Dashboard
ws_dash = wb.active
ws_dash.title = "Resumo Geral"
ws_dash.views.sheetView[0].showGridLines = True

# Aba principal de dados
ws = wb.create_sheet(title="Controle de Fretes")
ws.views.sheetView[0].showGridLines = True

# 2. Configuração de Estilos e Cores
NAVY = "1B365D"
TOTAL_BG = "E6EEF4"
ZEBRA_BG = "F8FAFC"
BORDER_GRAY = "D9D9D9"

font_title = Font(name="Segoe UI", size=16, bold=True, color=NAVY)
font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
font_regular = Font(name="Segoe UI", size=11)
font_bold = Font(name="Segoe UI", size=11, bold=True, color=NAVY)

fill_header = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
fill_zebra = PatternFill(start_color=ZEBRA_BG, end_color=ZEBRA_BG, fill_type="solid")
fill_total = PatternFill(start_color=TOTAL_BG, end_color=TOTAL_BG, fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color=BORDER_GRAY),
    right=Side(style='thin', color=BORDER_GRAY),
    top=Side(style='thin', color=BORDER_GRAY),
    bottom=Side(style='thin', color=BORDER_GRAY)
)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

# 3. Título e Cabeçalhos
ws.merge_cells("A1:J1")
ws["A1"] = "REGISTRO DE CONTROLE DE FRETES E LOGÍSTICA"
ws["A1"].font = font_title

headers = [
    "Data",
    "Fretista",
    "Obra",
    "Fornecedor",
    "Quantidade",
    "KM",
    "KM Total (Qtd x KM)",
    "Pedágio (R$)",
    "Valor Frete/KM (R$)",
    "Total Geral (R$)"
]

for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=3, column=col_idx, value=header)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = thin_border

# 4. Dados de Exemplo
dados_exemplo = [
    ("2026-10-01", "Carlos Silva", "Obra Residencial Alpha", "Pedreira Vale", 5, 45.0, 12.50, 8.00),
    ("2026-10-02", "João Pereira", "Obra Centro Comercial", "Cimento Campeão", 3, 120.0, 45.00, 10.50),
    ("2026-10-03", "Carlos Silva", "Obra Ponte Norte", "Aço Brasil", 8, 30.5, 22.00, 7.50),
    ("2026-10-04", "Marcos Souza", "Obra Residencial Alpha", "Pedreira Vale", 4, 50.0, 15.00, 8.00),
]

start_row = 4
for i, row_data in enumerate(dados_exemplo):
    r = start_row + i
    data, fretista, obra, forn, qtd, km, pedagio, valor_km = row_data
    
    ws.cell(row=r, column=1, value=data).alignment = align_center
    ws.cell(row=r, column=2, value=fretista).alignment = align_left
    ws.cell(row=r, column=3, value=obra).alignment = align_left
    ws.cell(row=r, column=4, value=forn).alignment = align_left
    
    ws.cell(row=r, column=5, value=qtd).alignment = align_right
    ws.cell(row=r, column=6, value=km).alignment = align_right
    
    # FÓRMULA 1: Multiplica Quantidade por KM (Coluna G)
    ws.cell(row=r, column=7, value=f"=E{r}*F{r}").alignment = align_right
    
    ws.cell(row=r, column=8, value=pedagio).alignment = align_right
    ws.cell(row=r, column=9, value=valor_km).alignment = align_right
    
    # FÓRMULA 2: Soma do Frete Total ((KM Total * Valor/KM) + Pedágio) (Coluna J)
    ws.cell(row=r, column=10, value=f"=(G{r}*I{r})+H{r}").alignment = align_right
    
    # Formatatação numérica
    ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"
    ws.cell(row=r, column=5).number_format = "#,##0"
    ws.cell(row=r, column=6).number_format = "#,##0.0"
    ws.cell(row=r, column=7).number_format = "#,##0.0"
    ws.cell(row=r, column=8).number_format = 'R$ #,##0.00'
    ws.cell(row=r, column=9).number_format = 'R$ #,##0.00'
    ws.cell(row=r, column=10).number_format = 'R$ #,##0.00'
    
    # Zebra striping e bordas
    for c in range(1, 11):
        cell = ws.cell(row=r, column=c)
        cell.font = font_regular
        cell.border = thin_border
        if i % 2 == 1:
            cell.fill = fill_zebra

# 5. Linha de Totalização
tot_row = start_row + len(dados_exemplo)
ws.cell(row=tot_row, column=1, value="TOTAL GERAL").font = font_bold
ws.cell(row=tot_row, column=5, value=f"=SUM(E{start_row}:E{tot_row-1})")
ws.cell(row=tot_row, column=6, value=f"=SUM(F{start_row}:F{tot_row-1})")
ws.cell(row=tot_row, column=7, value=f"=SUM(G{start_row}:G{tot_row-1})")
ws.cell(row=tot_row, column=8, value=f"=SUM(H{start_row}:H{tot_row-1})")
ws.cell(row=tot_row, column=10, value=f"=SUM(J{start_row}:J{tot_row-1})")

# Formatação da linha total
for c in range(1, 11):
    cell = ws.cell(row=tot_row, column=c)
    cell.font = font_bold
    cell.fill = fill_total
    cell.border = thin_border

# 6. Ajuste Automático da Largura das Colunas
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

# Salvar arquivo
wb.save("Controle_de_Fretes.xlsx")
